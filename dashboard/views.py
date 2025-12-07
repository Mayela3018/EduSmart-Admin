from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.db.models import Count
from django.utils import timezone
from datetime import timedelta, datetime
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.http import HttpResponse
import csv

# Para Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import Users, Tasks, Events, Habits


# ============================
#           LOGIN
# ============================
def dashboard_login(request):
    # Si ya está autenticado, redirigir al dashboard
    if request.user.is_authenticated:
        return redirect("/dashboard/")
    
    if request.method == "POST":
        email = request.POST.get("email")
        password = request.POST.get("password")

        # Validar que los campos no estén vacíos
        if not email or not password:
            return render(request, "dashboard/login.html", {
                "error": "Por favor completa todos los campos"
            })

        user = authenticate(request, username=email, password=password)

        if user is not None:
            login(request, user)
            messages.success(request, f"¡Bienvenido {user.username}!")
            return redirect("/dashboard/")
        else:
            return render(request, "dashboard/login.html", {
                "error": "Credenciales incorrectas. Verifica tu email y contraseña."
            })

    return render(request, "dashboard/login.html")


# ============================
#         LOGOUT
# ============================
def logout_view(request):
    logout(request)
    messages.info(request, "Has cerrado sesión correctamente")
    return redirect("/")


# ============================
#      DASHBOARD PRINCIPAL
# ============================
@login_required(login_url="/")
def admin_dashboard(request):
    try:
        today = timezone.now().date()
        week_ago = today - timedelta(days=7)

        total_users = Users.objects.count()
        total_tasks = Tasks.objects.count()
        total_events = Events.objects.count()
        total_habits = Habits.objects.count()

        completed_tasks = Tasks.objects.filter(completed=True).count()
        pending_tasks = total_tasks - completed_tasks

        completion_rate = round((completed_tasks / total_tasks) * 100, 1) if total_tasks else 0

        tasks_this_week = Tasks.objects.filter(created_at__date__gte=week_ago).count()
        events_this_week = Events.objects.filter(fecha__gte=week_ago, fecha__lte=today).count()
        new_users_week = Users.objects.filter(created_at__date__gte=week_ago).count()

        tasks_by_priority = {
            "ALTA": Tasks.objects.filter(prioridad="ALTA").count(),
            "MEDIA": Tasks.objects.filter(prioridad="MEDIA").count(),
            "BAJA": Tasks.objects.filter(prioridad="BAJA").count(),
        }

        events_by_category = {
            "CLASE": Events.objects.filter(categoria="CLASE").count(),
            "EXAMEN": Events.objects.filter(categoria="EXAMEN").count(),
            "PERSONAL": Events.objects.filter(categoria="PERSONAL").count(),
        }

        next_week = today + timedelta(days=7)

        upcoming_tasks = Tasks.objects.filter(
            completed=False,
            fecha_limite__range=(today, next_week)
        ).order_by("fecha_limite")[:5]

        upcoming_events = Events.objects.filter(
            fecha__gte=today
        ).order_by("fecha", "hora")[:5]

        active_users = Users.objects.annotate(
            task_count=Count("tasks", distinct=True)
        ).order_by("-task_count")[:5]

        active_habits = Habits.objects.filter(activo=True).count()
        inactive_habits = Habits.objects.filter(activo=False).count()

        context = {
            "total_users": total_users,
            "total_tasks": total_tasks,
            "total_events": total_events,
            "total_habits": total_habits,
            "completed_tasks": completed_tasks,
            "pending_tasks": pending_tasks,
            "completion_rate": completion_rate,
            "tasks_this_week": tasks_this_week,
            "events_this_week": events_this_week,
            "new_users_week": new_users_week,
            "tasks_by_priority": tasks_by_priority,
            "events_by_category": events_by_category,
            "upcoming_tasks": upcoming_tasks,
            "upcoming_events": upcoming_events,
            "active_users": active_users,
            "active_habits": active_habits,
            "inactive_habits": inactive_habits,
            "today": today,
        }

        return render(request, "dashboard/dashboard_home.html", context)
    
    except Exception as e:
        messages.error(request, f"Error al cargar el dashboard: {str(e)}")
        return render(request, "dashboard/dashboard_home.html", {})


# ============================
#         LISTADOS
# ============================
@login_required(login_url="/")
def dashboard_usuarios(request):
    try:
        usuarios = Users.objects.all().order_by("-created_at")
        return render(request, "dashboard/usuarios.html", {"usuarios": usuarios})
    except Exception as e:
        messages.error(request, f"Error al cargar usuarios: {str(e)}")
        return redirect("/dashboard/")


@login_required(login_url="/")
def dashboard_tareas(request):
    try:
        tareas = Tasks.objects.all().order_by("-created_at")
        return render(request, "dashboard/tareas.html", {"tareas": tareas})
    except Exception as e:
        messages.error(request, f"Error al cargar tareas: {str(e)}")
        return redirect("/dashboard/")


@login_required(login_url="/")
def dashboard_eventos(request):
    try:
        eventos = Events.objects.all().order_by("-fecha")
        return render(request, "dashboard/eventos.html", {"eventos": eventos})
    except Exception as e:
        messages.error(request, f"Error al cargar eventos: {str(e)}")
        return redirect("/dashboard/")


# ============================
#         FORMULARIOS
# ============================
@login_required(login_url="/")
def tarea_nueva(request):
    if request.method == "POST":
        try:
            titulo = request.POST.get("titulo")
            descripcion = request.POST.get("descripcion")
            prioridad = request.POST.get("prioridad")
            fecha_limite = request.POST.get("fecha_limite")

            Tasks.objects.create(
                titulo=titulo,
                descripcion=descripcion,
                prioridad=prioridad,
                fecha_limite=fecha_limite if fecha_limite else None,
                completed=False
            )
            messages.success(request, "✅ Tarea creada exitosamente")
            return redirect("/dashboard/tareas/")
        except Exception as e:
            messages.error(request, f"Error al crear tarea: {str(e)}")

    return render(request, "dashboard/formularios/nueva_tarea.html")


@login_required(login_url="/")
def evento_nuevo(request):
    if request.method == "POST":
        try:
            titulo = request.POST.get("titulo")
            descripcion = request.POST.get("descripcion")
            categoria = request.POST.get("categoria")
            fecha = request.POST.get("fecha")
            hora = request.POST.get("hora")

            Events.objects.create(
                titulo=titulo,
                descripcion=descripcion,
                categoria=categoria,
                fecha=fecha,
                hora=hora if hora else None
            )
            messages.success(request, "✅ Evento creado exitosamente")
            return redirect("/dashboard/eventos/")
        except Exception as e:
            messages.error(request, f"Error al crear evento: {str(e)}")

    return render(request, "dashboard/formularios/nuevo_evento.html")


@login_required(login_url="/")
def usuario_nuevo(request):
    if request.method == "POST":
        try:
            name = request.POST.get("name")
            email = request.POST.get("email")
            tipo = request.POST.get("tipo")

            Users.objects.create(
                name=name,
                email=email,
                tipo=tipo
            )
            messages.success(request, "✅ Usuario creado exitosamente")
            return redirect("/dashboard/usuarios/")
        except Exception as e:
            messages.error(request, f"Error al crear usuario: {str(e)}")

    return render(request, "dashboard/formularios/nuevo_usuario.html")


@login_required(login_url="/")
def habito_nuevo(request):
    if request.method == "POST":
        try:
            nombre = request.POST.get("nombre")
            tipo = request.POST.get("tipo")
            meta_diaria = request.POST.get("meta_diaria")
            activo = request.POST.get("activo") == "on"

            Habits.objects.create(
                nombre=nombre,
                tipo=tipo,
                meta_diaria=meta_diaria if meta_diaria else 1,
                activo=activo
            )
            messages.success(request, "✅ Hábito creado exitosamente")
            return redirect("/dashboard/")
        except Exception as e:
            messages.error(request, f"Error al crear hábito: {str(e)}")

    return render(request, "dashboard/formularios/nuevo_habito.html")


# ============================
#    EXPORTACIÓN A EXCEL
# ============================

@login_required(login_url="/")
def export_usuarios_excel(request):
    """Exportar usuarios a Excel"""
    try:
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Usuarios"
        
        # Estilos
        header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        
        # Headers
        headers = ['ID', 'Nombre', 'Email', 'Tipo', 'Fecha de Registro']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Datos
        usuarios = Users.objects.all().order_by('-created_at')
        for row, usuario in enumerate(usuarios, start=2):
            ws.cell(row=row, column=1, value=usuario.id)
            ws.cell(row=row, column=2, value=usuario.name)
            ws.cell(row=row, column=3, value=usuario.email)
            ws.cell(row=row, column=4, value=usuario.tipo)
            ws.cell(row=row, column=5, value=usuario.created_at.strftime('%d/%m/%Y %H:%M'))
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 20
        
        # Preparar respuesta
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'usuarios_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/usuarios/")


@login_required(login_url="/")
def export_tareas_excel(request):
    """Exportar tareas a Excel"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Tareas"
        
        header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        
        headers = ['ID', 'Título', 'Descripción', 'Prioridad', 'Estado', 'Fecha Límite', 'Creado']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        tareas = Tasks.objects.all().order_by('-created_at')
        for row, tarea in enumerate(tareas, start=2):
            ws.cell(row=row, column=1, value=tarea.id)
            ws.cell(row=row, column=2, value=tarea.titulo)
            ws.cell(row=row, column=3, value=tarea.descripcion or '')
            ws.cell(row=row, column=4, value=tarea.prioridad)
            ws.cell(row=row, column=5, value='Completada' if tarea.completed else 'Pendiente')
            ws.cell(row=row, column=6, value=tarea.fecha_limite.strftime('%d/%m/%Y') if tarea.fecha_limite else '')
            ws.cell(row=row, column=7, value=tarea.created_at.strftime('%d/%m/%Y %H:%M'))
        
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 15
        ws.column_dimensions['G'].width = 20
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'tareas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/tareas/")


@login_required(login_url="/")
def export_eventos_excel(request):
    """Exportar eventos a Excel"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Eventos"
        
        header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        
        headers = ['ID', 'Título', 'Descripción', 'Categoría', 'Fecha', 'Hora', 'Creado']
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        eventos = Events.objects.all().order_by('-fecha')
        for row, evento in enumerate(eventos, start=2):
            ws.cell(row=row, column=1, value=evento.id)
            ws.cell(row=row, column=2, value=evento.titulo)
            ws.cell(row=row, column=3, value=evento.descripcion or '')
            ws.cell(row=row, column=4, value=evento.categoria)
            ws.cell(row=row, column=5, value=evento.fecha.strftime('%d/%m/%Y'))
            ws.cell(row=row, column=6, value=evento.hora.strftime('%H:%M') if evento.hora else '')
            ws.cell(row=row, column=7, value=evento.created_at.strftime('%d/%m/%Y %H:%M'))
        
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 20
        
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f'eventos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        wb.save(response)
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/eventos/")


# ============================
#    EXPORTACIÓN A CSV
# ============================

@login_required(login_url="/")
def export_usuarios_csv(request):
    """Exportar usuarios a CSV"""
    try:
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        filename = f'usuarios_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        # BOM para que Excel reconozca UTF-8
        response.write('\ufeff')
        
        writer = csv.writer(response)
        writer.writerow(['ID', 'Nombre', 'Email', 'Tipo', 'Fecha de Registro'])
        
        usuarios = Users.objects.all().order_by('-created_at')
        for usuario in usuarios:
            writer.writerow([
                usuario.id,
                usuario.name,
                usuario.email,
                usuario.tipo,
                usuario.created_at.strftime('%d/%m/%Y %H:%M')
            ])
        
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/usuarios/")


@login_required(login_url="/")
def export_tareas_csv(request):
    """Exportar tareas a CSV"""
    try:
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        filename = f'tareas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        response.write('\ufeff')
        
        writer = csv.writer(response)
        writer.writerow(['ID', 'Título', 'Descripción', 'Prioridad', 'Estado', 'Fecha Límite', 'Creado'])
        
        tareas = Tasks.objects.all().order_by('-created_at')
        for tarea in tareas:
            writer.writerow([
                tarea.id,
                tarea.titulo,
                tarea.descripcion or '',
                tarea.prioridad,
                'Completada' if tarea.completed else 'Pendiente',
                tarea.fecha_limite.strftime('%d/%m/%Y') if tarea.fecha_limite else '',
                tarea.created_at.strftime('%d/%m/%Y %H:%M')
            ])
        
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/tareas/")


@login_required(login_url="/")
def export_eventos_csv(request):
    """Exportar eventos a CSV"""
    try:
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        filename = f'eventos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        response.write('\ufeff')
        
        writer = csv.writer(response)
        writer.writerow(['ID', 'Título', 'Descripción', 'Categoría', 'Fecha', 'Hora', 'Creado'])
        
        eventos = Events.objects.all().order_by('-fecha')
        for evento in eventos:
            writer.writerow([
                evento.id,
                evento.titulo,
                evento.descripcion or '',
                evento.categoria,
                evento.fecha.strftime('%d/%m/%Y'),
                evento.hora.strftime('%H:%M') if evento.hora else '',
                evento.created_at.strftime('%d/%m/%Y %H:%M')
            ])
        
        return response
        
    except Exception as e:
        messages.error(request, f"Error al exportar: {str(e)}")
        return redirect("/dashboard/eventos/")